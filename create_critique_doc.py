from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_paragraph('Critique Recommendation Questions')
title_run = title.runs[0]
title_run.font.size = Pt(24)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0, 0, 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add some space
doc.add_paragraph()

# List of recommendations
recommendations = [
    'Remove the point of view for accuracy of writing',
    'Frontend in admin, the colors should be changed, some text are not a bit readable.',
    'Include options in choosing the style of writing.',
    'It would better if the user can input multiple excerpts for better style recognition/model training.',
    'Introduce multiple options for the generated text on which it can be used for feedback collection and analytics.',
    'Can it filter out unnecessary words? (e.g inappropriate, vulgar tones)',
    'Hints and helpful information on what the tuning parameter are for and how it impacts the paraphrase',
    'Report button',
    'Flagging of inappropriate words? (deflammatory words)',
    'Context checking',
    'Revisit Business model',
    '"Profile" -> "Style"',
    'UI Improvements: Icons, Consistency, Colors',
    'Accuracy of results',
    'Suggestions for writing improvements'
]

for i, rec in enumerate(recommendations, 1):
    p = doc.add_paragraph(f'{i}. {rec}')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)

# Save the document
output_path = r'C:\Users\MSI\Documents\GitHub\new_Stylesync\Panelist_Critique_Recommendations.docx'
doc.save(output_path)
print(f'Document created: {output_path}')
